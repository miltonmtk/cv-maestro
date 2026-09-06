from __future__ import annotations

"""
EXPORTADOR FINAL POR LOTE
Proyecto: cv-maestro

Genera automáticamente para cada vacante:

    DOCX + PDF

Modos:
    SIN FOTO:
        python exportador_final_lote.py

    CON FOTO:
        python exportador_final_lote.py --foto foto_cv.jpg

Reglas:
- No modifica CV_MAESTRO.
- Cada vacante se procesa independientemente.
- Solo exporta CV con auditoría APROBADA.
- La fotografía es opcional.
- La fotografía no se convierte en evidencia profesional.
- El cuerpo del CV permanece en una sola columna y ATS-friendly.
"""

import argparse
import shutil
import subprocess
from pathlib import Path
from typing import Any, Dict, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from cargador_vacantes import cargar_vacante
from generador_cv import (
    auditar_cv,
    construir_cv,
    renderizar_cv,
    slug,
)

from exportador_documentos import (
    agregar_cabecera,
    agregar_competencias,
    agregar_experiencias,
    agregar_formacion,
    agregar_perfil,
    configurar_documento,
    configurar_metadatos,
)


# ============================================================
# CONFIGURACION
# ============================================================

CARPETA_VACANTES = Path("vacantes")

CARPETA_FINAL = (
    Path("salidas")
    / "documentos_finales"
)

SOFFICE_WINDOWS = Path(
    r"C:\Program Files\LibreOffice\program\soffice.exe"
)

EXTENSIONES_FOTO = {
    ".jpg",
    ".jpeg",
    ".png",
}


# ============================================================
# LIBREOFFICE
# ============================================================

def localizar_soffice() -> Path:

    if SOFFICE_WINDOWS.exists():
        return SOFFICE_WINDOWS

    encontrado = shutil.which(
        "soffice"
    )

    if encontrado:
        return Path(encontrado)

    raise FileNotFoundError(
        "No se encontró LibreOffice/soffice."
    )


# ============================================================
# FOTOGRAFIA
# ============================================================

def validar_foto(
    ruta_foto: Optional[Path],
) -> Optional[Path]:

    if ruta_foto is None:
        return None

    ruta_foto = Path(
        ruta_foto
    )

    if not ruta_foto.exists():
        raise FileNotFoundError(
            f"No existe la fotografía: "
            f"{ruta_foto}"
        )

    if (
        ruta_foto.suffix.lower()
        not in EXTENSIONES_FOTO
    ):
        raise ValueError(
            "La fotografía debe ser "
            "JPG, JPEG o PNG."
        )

    return ruta_foto


def agregar_fotografia(
    documento: Document,
    ruta_foto: Path,
) -> None:

    """
    Coloca la fotografía en la cabecera,
    alineada a la derecha.

    El cuerpo profesional permanece
    en una sola columna.
    """

    seccion = documento.sections[0]

    # Dejamos espacio suficiente para
    # evitar que la fotografía invada
    # el contenido principal.
    seccion.top_margin = Cm(3.3)
    seccion.header_distance = Cm(0.35)

    cabecera = seccion.header

    parrafo = cabecera.paragraphs[0]

    parrafo.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    parrafo.paragraph_format.space_after = 0

    run = parrafo.add_run()

    run.add_picture(
        str(ruta_foto),
        width=Cm(2.7),
    )


# ============================================================
# DOCX
# ============================================================

def crear_docx(
    cv: Dict[str, Any],
    ruta_docx: Path,
    ruta_foto: Optional[Path],
) -> Path:

    documento = Document()

    configurar_documento(
        documento
    )

    configurar_metadatos(
        documento,
        cv,
    )

    if ruta_foto is not None:
        agregar_fotografia(
            documento,
            ruta_foto,
        )

    agregar_cabecera(
        documento,
        cv,
    )

    agregar_perfil(
        documento,
        cv,
    )

    agregar_competencias(
        documento,
        cv,
    )

    agregar_experiencias(
        documento,
        cv,
    )

    agregar_formacion(
        documento,
        cv,
    )

    ruta_docx.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    documento.save(
        ruta_docx
    )

    return ruta_docx


# ============================================================
# PDF
# ============================================================

def convertir_pdf(
    ruta_docx: Path,
) -> Path:

    soffice = localizar_soffice()

    ruta_pdf = (
        ruta_docx.with_suffix(
            ".pdf"
        )
    )

    # Eliminamos una versión anterior
    # para no confundir una conversión
    # nueva con un archivo viejo.
    if ruta_pdf.exists():
        ruta_pdf.unlink()

    comando = [
        str(soffice),
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(
            ruta_docx.parent
        ),
        str(ruta_docx),
    ]

    resultado = subprocess.run(
        comando,
        capture_output=True,
        text=True,
        timeout=60,
    )

    if (
        resultado.returncode != 0
        or not ruta_pdf.exists()
    ):

        detalle = (
            resultado.stderr.strip()
            or resultado.stdout.strip()
            or "LibreOffice no generó el PDF."
        )

        raise RuntimeError(
            "Error convirtiendo DOCX a PDF: "
            + detalle
        )

    return ruta_pdf


# ============================================================
# NOMBRE DE SALIDA
# ============================================================

def nombre_base(
    vacante: Dict[str, Any],
    archivo_origen: Path,
    con_foto: bool,
) -> str:

    empresa = slug(
        vacante.get(
            "empresa",
            "EMPRESA",
        )
    )

    origen = slug(
        archivo_origen.stem
    )

    modo = (
        "CON_FOTO"
        if con_foto
        else "SIN_FOTO"
    )

    return (
        f"CV_{empresa}_"
        f"{origen}_"
        f"M_MONTANO_"
        f"{modo}"
    )


# ============================================================
# PROCESAR UNA VACANTE
# ============================================================

def procesar_vacante(
    archivo: Path,
    ruta_foto: Optional[Path],
) -> Dict[str, Any]:

    vacante = cargar_vacante(
        archivo
    )

    cv = construir_cv(
        vacante
    )

    texto_cv = renderizar_cv(
        cv
    )

    auditoria = auditar_cv(
        cv,
        texto_cv,
    )

    if (
        auditoria.get(
            "estado"
        )
        != "APROBADO"
    ):

        return {
            "archivo": str(
                archivo
            ),
            "empresa": vacante.get(
                "empresa",
                "",
            ),
            "titulo": vacante.get(
                "titulo",
                "",
            ),
            "estado": "BLOQUEADA",
            "docx": None,
            "pdf": None,
            "incidencias": auditoria.get(
                "incidencias",
                [],
            ),
        }

    con_foto = (
        ruta_foto is not None
    )

    modo_carpeta = (
        "con_foto"
        if con_foto
        else "sin_foto"
    )

    carpeta = (
        CARPETA_FINAL
        / modo_carpeta
    )

    base = nombre_base(
        vacante,
        archivo,
        con_foto,
    )

    ruta_docx = (
        carpeta
        / f"{base}.docx"
    )

    crear_docx(
        cv,
        ruta_docx,
        ruta_foto,
    )

    ruta_pdf = convertir_pdf(
        ruta_docx
    )

    return {
        "archivo": str(
            archivo
        ),
        "empresa": vacante.get(
            "empresa",
            "",
        ),
        "titulo": vacante.get(
            "titulo",
            "",
        ),
        "estado": "GENERADA",
        "docx": str(
            ruta_docx
        ),
        "pdf": str(
            ruta_pdf
        ),
        "incidencias": [],
    }


# ============================================================
# PROCESAR TODAS LAS VACANTES
# ============================================================

def procesar_lote(
    ruta_foto: Optional[Path],
) -> list[Dict[str, Any]]:

    if not CARPETA_VACANTES.exists():

        raise FileNotFoundError(
            "No existe la carpeta vacantes."
        )

    archivos = sorted(
        CARPETA_VACANTES.glob(
            "*.json"
        )
    )

    if not archivos:

        raise FileNotFoundError(
            "No hay archivos JSON "
            "en la carpeta vacantes."
        )

    resultados = []

    for archivo in archivos:

        try:

            resultado = procesar_vacante(
                archivo,
                ruta_foto,
            )

        except Exception as error:

            resultado = {
                "archivo": str(
                    archivo
                ),
                "empresa": "",
                "titulo": "",
                "estado": "ERROR",
                "docx": None,
                "pdf": None,
                "incidencias": [
                    str(error)
                ],
            }

        resultados.append(
            resultado
        )

    return resultados


# ============================================================
# MOSTRAR RESULTADOS
# ============================================================

def mostrar_resultados(
    resultados: list[
        Dict[str, Any]
    ],
    con_foto: bool,
) -> None:

    print()
    print("=" * 72)
    print("EXPORTADOR FINAL DOCX + PDF")
    print("=" * 72)

    print(
        "Modo:",
        (
            "CON FOTOGRAFIA"
            if con_foto
            else "SIN FOTOGRAFIA"
        ),
    )

    print()

    for indice, resultado in enumerate(
        resultados,
        start=1,
    ):

        print(
            f"{indice}. "
            f"{resultado.get('titulo') or 'VACANTE'}"
        )

        if resultado.get(
            "empresa"
        ):

            print(
                "   Empresa:",
                resultado[
                    "empresa"
                ],
            )

        print(
            "   Estado:",
            resultado[
                "estado"
            ],
        )

        if resultado.get(
            "docx"
        ):

            print(
                "   DOCX:",
                resultado[
                    "docx"
                ],
            )

        if resultado.get(
            "pdf"
        ):

            print(
                "   PDF:",
                resultado[
                    "pdf"
                ],
            )

        for incidencia in resultado.get(
            "incidencias",
            [],
        ):

            print(
                "   Incidencia:",
                incidencia,
            )

        print()

    generadas = sum(
        1
        for resultado in resultados
        if resultado[
            "estado"
        ]
        == "GENERADA"
    )

    errores = sum(
        1
        for resultado in resultados
        if resultado[
            "estado"
        ]
        == "ERROR"
    )

    bloqueadas = sum(
        1
        for resultado in resultados
        if resultado[
            "estado"
        ]
        == "BLOQUEADA"
    )

    print("=" * 72)

    print(
        "VACANTES:",
        len(resultados),
    )

    print(
        "DOCX + PDF GENERADOS:",
        generadas,
    )

    print(
        "BLOQUEADAS:",
        bloqueadas,
    )

    print(
        "ERRORES:",
        errores,
    )

    print("=" * 72)


# ============================================================
# ARGUMENTOS
# ============================================================

def leer_argumentos():

    parser = argparse.ArgumentParser(
        description=(
            "Genera DOCX y PDF para "
            "todas las vacantes."
        )
    )

    parser.add_argument(
        "--foto",
        type=Path,
        default=None,
        help=(
            "Ruta opcional a fotografía "
            "JPG, JPEG o PNG."
        ),
    )

    return parser.parse_args()


# ============================================================
# EJECUCION
# ============================================================

def main() -> None:

    argumentos = leer_argumentos()

    ruta_foto = validar_foto(
        argumentos.foto
    )

    resultados = procesar_lote(
        ruta_foto
    )

    mostrar_resultados(
        resultados,
        ruta_foto is not None,
    )


if __name__ == "__main__":
    main()