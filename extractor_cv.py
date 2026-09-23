from __future__ import annotations

import re
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from pypdf import PdfReader


MAX_CV_BYTES = 10 * 1024 * 1024
MAX_CV_PAGINAS = 40
MAX_TEXTO_EXTRAIDO = 150_000
EXTENSIONES_CV = {".pdf", ".docx"}


class ExtractorCVError(ValueError):
    """Error controlado durante la lectura de un CV."""


def _limpiar_texto(texto: str) -> str:
    lineas = []
    for linea in texto.replace("\x00", "").splitlines():
        limpia = re.sub(r"[ \t]+", " ", linea).strip()
        if limpia:
            lineas.append(limpia)
    return "\n".join(lineas)[:MAX_TEXTO_EXTRAIDO]


def _validar_archivo(contenido: bytes, nombre: str) -> str:
    if not contenido:
        raise ExtractorCVError("El CV está vacío.")
    if len(contenido) > MAX_CV_BYTES:
        raise ExtractorCVError("El CV supera el límite de 10 MB.")
    extension = Path(nombre).suffix.lower()
    if extension not in EXTENSIONES_CV:
        raise ExtractorCVError("El CV debe estar en formato PDF o DOCX.")
    return extension


def _extraer_pdf(contenido: bytes) -> tuple[str, int]:
    try:
        lector = PdfReader(BytesIO(contenido))
        if lector.is_encrypted:
            try:
                desbloqueado = lector.decrypt("")
            except Exception as error:
                raise ExtractorCVError(
                    "El PDF está protegido y no puede procesarse."
                ) from error
            if not desbloqueado:
                raise ExtractorCVError(
                    "El PDF está protegido y no puede procesarse."
                )
        if len(lector.pages) > MAX_CV_PAGINAS:
            raise ExtractorCVError(
                f"El PDF supera el límite de {MAX_CV_PAGINAS} páginas."
            )
        paginas = [pagina.extract_text() or "" for pagina in lector.pages]
    except ExtractorCVError:
        raise
    except Exception as error:
        raise ExtractorCVError(
            "No fue posible leer el PDF. Verifica que no esté dañado."
        ) from error
    return _limpiar_texto("\n".join(paginas)), len(paginas)


def _extraer_docx(contenido: bytes) -> tuple[str, int]:
    try:
        documento = Document(BytesIO(contenido))
        bloques: List[str] = [
            parrafo.text for parrafo in documento.paragraphs if parrafo.text.strip()
        ]
        for tabla in documento.tables:
            for fila in tabla.rows:
                celdas = [celda.text.strip() for celda in fila.cells if celda.text.strip()]
                if celdas:
                    bloques.append(" | ".join(celdas))
    except Exception as error:
        raise ExtractorCVError(
            "No fue posible leer el DOCX. Verifica que no esté dañado."
        ) from error
    return _limpiar_texto("\n".join(bloques)), 1


def extraer_texto_cv(contenido: bytes, nombre: str) -> Dict[str, Any]:
    extension = _validar_archivo(contenido, nombre)
    if extension == ".pdf":
        texto, paginas = _extraer_pdf(contenido)
    else:
        texto, paginas = _extraer_docx(contenido)

    if len(texto) < 40:
        mensaje = (
            "El CV no contiene texto extraíble. "
            "Si es un PDF escaneado, conviértelo con OCR antes de cargarlo."
        )
        raise ExtractorCVError(mensaje)

    return {
        "nombre_archivo": Path(nombre).name,
        "extension": extension,
        "paginas": paginas,
        "texto": texto,
        "caracteres": len(texto),
    }


def _primera_coincidencia(patron: str, texto: str) -> str:
    coincidencia = re.search(patron, texto, flags=re.IGNORECASE)
    return coincidencia.group(0).strip() if coincidencia else ""


def _nombre_probable(lineas: List[str]) -> str:
    prohibidas = {
        "curriculum vitae",
        "currículum vitae",
        "cv",
        "perfil profesional",
        "experiencia profesional",
    }
    for indice, linea in enumerate(lineas[:12]):
        candidata = linea.strip(" |-")
        palabras = candidata.split()
        if (
            2 <= len(palabras) <= 7
            and candidata.lower() not in prohibidas
            and not re.search(r"[@\d:/]", candidata)
            and len(candidata) <= 100
        ):
            partes = [candidata]
            for siguiente in lineas[indice + 1:indice + 3]:
                siguiente = siguiente.strip(" |-")
                if (
                    1 <= len(siguiente.split()) <= 3
                    and siguiente == siguiente.upper()
                    and re.fullmatch(
                        r"[A-ZÁÉÍÓÚÜÑ][A-ZÁÉÍÓÚÜÑ '\u2019-]*",
                        siguiente,
                    )
                    and not re.search(r"[@\d:/]", siguiente)
                    and not _clave_encabezado(siguiente)
                ):
                    partes.append(siguiente)
                else:
                    break
            return " ".join(partes)[:100]
    return ""


def _sin_acentos(texto: str) -> str:
    return "".join(
        caracter
        for caracter in unicodedata.normalize("NFD", texto)
        if unicodedata.category(caracter) != "Mn"
    )


def _ubicacion_probable(lineas: List[str], email: str, telefono: str) -> str:
    """Obtiene solo una ubicación explícita del bloque inicial del CV."""
    for linea in lineas[:15]:
        if email and email not in linea and telefono and telefono not in linea:
            continue
        segmentos = [segmento.strip(" |-·") for segmento in re.split(r"[|•]", linea)]
        for segmento in segmentos:
            if not segmento or "@" in segmento or re.search(r"\d{5,}", segmento):
                continue
            normalizado = _sin_acentos(segmento.lower())
            if "linkedin" in normalizado or len(segmento.split()) > 6:
                continue
            if re.fullmatch(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ .,'’-]+", segmento):
                return segmento
    return ""


PATRON_PERIODO = (
    r"(?:19|20)\d{2}\s*(?:[-–—/]|\ba\b)\s*"
    r"(?:actualidad|presente|(?:19|20)\d{2})"
)

PATRON_CABECERA_EXPERIENCIA = re.compile(
    r"(?<![\wÁÉÍÓÚÜÑ])"
    r"(?P<organizacion>"
    r"[A-ZÁÉÍÓÚÜÑ][A-ZÁÉÍÓÚÜÑ0-9&.'’/-]*"
    r"(?:\s+(?:DE|DEL|LA|LAS|LOS|Y|EN|[A-ZÁÉÍÓÚÜÑ0-9&.'’/-]+)){1,8}"
    r")\s+"
    r"(?P<cargo>[^|\n]{2,80}?)\s*\|\s*"
    rf"(?P<periodo>{PATRON_PERIODO})",
)


def _frases(texto: str) -> List[str]:
    return [
        frase.strip(" .•-\n\t")
        for frase in re.split(r"(?:\.\s+|[•\n]+)", texto)
        if frase.strip(" .•-\n\t")
    ]


def _proponer_experiencias(texto: str) -> List[Dict[str, Any]]:
    coincidencias = list(PATRON_CABECERA_EXPERIENCIA.finditer(texto))
    propuestas: List[Dict[str, Any]] = []
    for indice, coincidencia in enumerate(coincidencias):
        fin = coincidencias[indice + 1].start() if indice + 1 < len(coincidencias) else len(texto)
        propuestas.append({
            "organizacion": coincidencia.group("organizacion").strip(),
            "cargo": coincidencia.group("cargo").strip(" .·-"),
            "periodo": coincidencia.group("periodo").strip(),
            "area": [],
            "funciones": _frases(texto[coincidencia.end():fin]),
        })
    return propuestas


def _proponer_formacion(texto: str) -> List[Dict[str, str]]:
    """Propone estudios solo cuando existen separadores explícitos; no adivina."""
    propuestas: List[Dict[str, str]] = []
    for linea in texto.splitlines():
        partes = [parte.strip(" .•-") for parte in re.split(r"\s*[|]\s*", linea)]
        if len(partes) < 2:
            continue
        periodo = ""
        if re.fullmatch(r"(?:19|20)\d{2}(?:\s*[-–—/]\s*(?:19|20)\d{2})?", partes[-1]):
            periodo = partes.pop()
        if len(partes) >= 2:
            propuestas.append({
                "titulo": partes[0],
                "institucion": partes[1],
                "periodo": periodo,
            })
    return propuestas


ENCABEZADOS = {
    "perfil": (
        "perfil profesional", "perfil", "resumen profesional",
        "sobre mí", "sobre mi", "objetivo profesional",
    ),
    "experiencia": (
        "experiencia profesional", "experiencia laboral", "experiencia",
        "trayectoria profesional",
    ),
    "formacion": (
        "formación académica", "formacion academica", "formación",
        "formacion", "educación", "educacion", "estudios",
    ),
    "competencias": (
        "competencias", "habilidades", "aptitudes", "skills",
        "conocimientos técnicos", "conocimientos tecnicos",
    ),
}


def _clave_encabezado(linea: str) -> str:
    normalizada = re.sub(r"[^a-záéíóúüñ ]", "", linea.lower()).strip()
    for clave, nombres in ENCABEZADOS.items():
        if normalizada in nombres:
            return clave
    return ""


def _secciones(texto: str) -> Dict[str, List[str]]:
    resultado = {clave: [] for clave in ENCABEZADOS}
    actual = ""
    for linea in texto.splitlines():
        clave = _clave_encabezado(linea)
        if clave:
            actual = clave
        elif actual:
            resultado[actual].append(linea)
    return resultado


def proponer_borrador(texto: str) -> Dict[str, Any]:
    texto = _limpiar_texto(texto)
    lineas = texto.splitlines()
    secciones = _secciones(texto)

    email = _primera_coincidencia(
        r"(?<![\w.+-])[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}(?![\w.-])",
        texto,
    )
    linkedin = _primera_coincidencia(
        r"(?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9_%?=./-]+",
        texto,
    )
    telefono = _primera_coincidencia(
        r"(?<!\d)(?:\+?\d{1,3}[ .-]?)?(?:\(?\d{2,4}\)?[ .-]?)"
        r"\d{3}[ .-]?\d{3,4}(?!\d)",
        texto,
    )

    perfil = "\n".join(secciones["perfil"][:8]).strip()
    competencias_lineas = secciones["competencias"][:20]
    competencias: List[str] = []
    for linea in competencias_lineas:
        for parte in re.split(r"[|•;,]", linea):
            limpia = parte.strip(" -\t")
            if limpia and limpia not in competencias:
                competencias.append(limpia)

    texto_formacion = "\n".join(secciones["formacion"]).strip()
    texto_experiencia = "\n".join(secciones["experiencia"]).strip()

    return {
        "datos_personales": {
            "nombre": _nombre_probable(lineas),
            "ubicacion": _ubicacion_probable(lineas, email, telefono),
            "telefono": telefono,
            "email": email,
            "linkedin": linkedin,
        },
        "perfil_profesional": perfil,
        "competencias": competencias,
        "formacion": _proponer_formacion(texto_formacion),
        "experiencia": _proponer_experiencias(texto_experiencia),
        "texto_formacion": texto_formacion,
        "texto_experiencia": texto_experiencia,
        "texto_fuente": texto,
    }
