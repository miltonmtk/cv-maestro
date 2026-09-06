from __future__ import annotations

"""
EXPORTADOR DE DOCUMENTOS ATS
Proyecto: cv-maestro

Objetivo:
- Recibir el CV estructurado generado por generador_cv.py.
- Crear un DOCX profesional y ATS-friendly.
- Una sola columna.
- Sin tablas.
- Sin iconos.
- Sin barras, gráficos ni elementos decorativos.
- No modifica CV_MAESTRO.
- No incorpora brechas, IAP ni controles internos al CV visible.
"""

from pathlib import Path
from typing import Any, Dict, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from cargador_vacantes import cargar_vacante
from generador_cv import (
    auditar_cv,
    construir_cv,
    renderizar_cv,
    slug,
)


CARPETA_DOCUMENTOS = Path("salidas") / "documentos"

FUENTE = "Arial"
TAMANO_TEXTO = 10.5
TAMANO_NOMBRE = 16
TAMANO_TITULO = 11
TAMANO_SECCION = 11

MARGEN_SUPERIOR = 1.5
MARGEN_INFERIOR = 1.5
MARGEN_IZQUIERDO = 1.7
MARGEN_DERECHO = 1.7


# ============================================================
# UTILIDADES
# ============================================================

def limpiar(valor: Any) -> str:
    return " ".join(
        str(valor or "").split()
    ).strip()


def campo(
    registro: Dict[str, Any],
    *nombres: str,
) -> Any:

    objetivos = {
        nombre.lower().strip()
        for nombre in nombres
    }

    for clave, valor in registro.items():
        if str(clave).lower().strip() in objetivos:
            return valor

    return ""


def lista_limpia(
    valores: Iterable[Any],
) -> list[str]:

    salida = []

    for valor in valores:
        texto = limpiar(valor)

        if texto:
            salida.append(texto)

    return salida


# ============================================================
# CONFIGURACION DEL DOCUMENTO
# ============================================================

def configurar_documento(
    documento: Document,
) -> None:

    seccion = documento.sections[0]

    seccion.top_margin = Cm(
        MARGEN_SUPERIOR
    )

    seccion.bottom_margin = Cm(
        MARGEN_INFERIOR
    )

    seccion.left_margin = Cm(
        MARGEN_IZQUIERDO
    )

    seccion.right_margin = Cm(
        MARGEN_DERECHO
    )

    seccion.header_distance = Cm(0.5)
    seccion.footer_distance = Cm(0.5)

    estilo_normal = documento.styles[
        "Normal"
    ]

    estilo_normal.font.name = FUENTE
    estilo_normal.font.size = Pt(
        TAMANO_TEXTO
    )

    estilo_normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        FUENTE,
    )

    formato = estilo_normal.paragraph_format

    formato.space_after = Pt(2)
    formato.space_before = Pt(0)
    formato.line_spacing = 1.0


# ============================================================
# PARRAFOS
# ============================================================

def agregar_parrafo(
    documento: Document,
    texto: str = "",
    *,
    negrita: bool = False,
    tamano: float | None = None,
    espacio_antes: float = 0,
    espacio_despues: float = 2,
) -> None:

    parrafo = documento.add_paragraph()

    parrafo.paragraph_format.space_before = Pt(
        espacio_antes
    )

    parrafo.paragraph_format.space_after = Pt(
        espacio_despues
    )

    parrafo.paragraph_format.line_spacing = 1.0

    run = parrafo.add_run(
        limpiar(texto)
    )

    run.bold = negrita
    run.font.name = FUENTE
    run.font.size = Pt(
        tamano or TAMANO_TEXTO
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        FUENTE,
    )


def agregar_seccion(
    documento: Document,
    titulo: str,
) -> None:

    parrafo = documento.add_paragraph()

    parrafo.paragraph_format.space_before = Pt(7)
    parrafo.paragraph_format.space_after = Pt(3)
    parrafo.paragraph_format.keep_with_next = True

    run = parrafo.add_run(
        limpiar(titulo).upper()
    )

    run.bold = True
    run.font.name = FUENTE
    run.font.size = Pt(
        TAMANO_SECCION
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        FUENTE,
    )


def agregar_vineta(
    documento: Document,
    texto: str,
) -> None:

    texto = limpiar(texto)

    if not texto:
        return

    parrafo = documento.add_paragraph(
        style=None
    )

    parrafo.paragraph_format.left_indent = Cm(
        0.35
    )

    parrafo.paragraph_format.first_line_indent = Cm(
        -0.25
    )

    parrafo.paragraph_format.space_before = Pt(0)
    parrafo.paragraph_format.space_after = Pt(1.5)
    parrafo.paragraph_format.line_spacing = 1.0

    run = parrafo.add_run(
        "• " + texto
    )

    run.font.name = FUENTE
    run.font.size = Pt(
        TAMANO_TEXTO
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        FUENTE,
    )


# ============================================================
# CABECERA DEL CV
# ============================================================

def agregar_cabecera(
    documento: Document,
    cv: Dict[str, Any],
) -> None:

    identidad = cv.get(
        "identidad",
        {},
    )

    nombre = limpiar(
        identidad.get(
            "nombre",
            "",
        )
    )

    if nombre:
        agregar_parrafo(
            documento,
            nombre.upper(),
            negrita=True,
            tamano=TAMANO_NOMBRE,
            espacio_despues=1,
        )

    contacto = lista_limpia(
        (
            identidad.get(
                "ubicacion",
                "",
            ),
            identidad.get(
                "telefono",
                "",
            ),
            identidad.get(
                "email",
                "",
            ),
            identidad.get(
                "linkedin",
                "",
            ),
        )
    )

    if contacto:
        agregar_parrafo(
            documento,
            " | ".join(contacto),
            tamano=9.5,
            espacio_despues=5,
        )

    titulo = limpiar(
        cv.get(
            "vacante",
            {},
        ).get(
            "titulo",
            "",
        )
    )

    if titulo:
        agregar_parrafo(
            documento,
            titulo,
            negrita=True,
            tamano=TAMANO_TITULO,
            espacio_despues=5,
        )


# ============================================================
# PERFIL PROFESIONAL
# ============================================================

def agregar_perfil(
    documento: Document,
    cv: Dict[str, Any],
) -> None:

    perfil = limpiar(
        cv.get(
            "perfil_profesional",
            "",
        )
    )

    if not perfil:
        return

    agregar_seccion(
        documento,
        "Perfil profesional",
    )

    agregar_parrafo(
        documento,
        perfil,
        espacio_despues=3,
    )


# ============================================================
# COMPETENCIAS
# ============================================================

def agregar_competencias(
    documento: Document,
    cv: Dict[str, Any],
) -> None:

    competencias = lista_limpia(
        cv.get(
            "competencias",
            [],
        )
    )

    if not competencias:
        return

    agregar_seccion(
        documento,
        "Competencias relevantes",
    )

    agregar_parrafo(
        documento,
        " | ".join(
            competencias
        ),
        espacio_despues=3,
    )


# ============================================================
# EXPERIENCIA
# ============================================================

def agregar_experiencias(
    documento: Document,
    cv: Dict[str, Any],
) -> None:

    experiencias = cv.get(
        "experiencias",
        [],
    )

    if not experiencias:
        return

    agregar_seccion(
        documento,
        "Experiencia profesional relevante",
    )

    for experiencia in experiencias:

        if not isinstance(
            experiencia,
            dict,
        ):
            continue

        cargo = limpiar(
            campo(
                experiencia,
                "cargo",
                "puesto",
            )
        )

        organizacion = limpiar(
            campo(
                experiencia,
                "organizacion",
                "organización",
                "empresa",
                "institucion",
            )
        )

        periodo = limpiar(
            campo(
                experiencia,
                "periodo",
                "período",
                "fecha",
            )
        )

        cabecera = " | ".join(
            valor
            for valor in (
                cargo,
                organizacion,
                periodo,
            )
            if valor
        )

        if cabecera:
            agregar_parrafo(
                documento,
                cabecera,
                negrita=True,
                espacio_antes=3,
                espacio_despues=1,
            )

        funciones = experiencia.get(
            "funciones_relevantes",
            [],
        )

        for funcion in funciones:
            agregar_vineta(
                documento,
                funcion,
            )


# ============================================================
# FORMACION
# ============================================================

def texto_formacion(
    elemento: Any,
) -> str:

    if not isinstance(
        elemento,
        dict,
    ):
        return limpiar(
            elemento
        )

    titulo = limpiar(
        campo(
            elemento,
            "titulo",
            "título",
            "nombre",
            "curso",
        )
    )

    institucion = limpiar(
        campo(
            elemento,
            "institucion",
            "institución",
            "entidad",
            "organizacion",
        )
    )

    fecha = limpiar(
        campo(
            elemento,
            "fecha",
            "periodo",
            "año",
            "ano",
            "anio",
        )
    )

    horas = limpiar(
        campo(
            elemento,
            "horas",
            "duracion",
            "duración",
        )
    )

    return " | ".join(
        valor
        for valor in (
            titulo,
            institucion,
            fecha,
            horas,
        )
        if valor
    )


def agregar_formacion(
    documento: Document,
    cv: Dict[str, Any],
) -> None:

    formacion = cv.get(
        "formacion",
        [],
    )

    if not formacion:
        return

    agregar_seccion(
        documento,
        "Formacion relevante",
    )

    for elemento in formacion:

        texto = texto_formacion(
            elemento
        )

        if texto:
            agregar_vineta(
                documento,
                texto,
            )


# ============================================================
# METADATOS DEL DOCX
# ============================================================

def configurar_metadatos(
    documento: Document,
    cv: Dict[str, Any],
) -> None:

    propiedades = documento.core_properties

    identidad = cv.get(
        "identidad",
        {},
    )

    vacante = cv.get(
        "vacante",
        {},
    )

    propiedades.author = limpiar(
        identidad.get(
            "nombre",
            "",
        )
    )

    propiedades.title = (
        "CV - "
        + limpiar(
            vacante.get(
                "titulo",
                "",
            )
        )
    )

    propiedades.subject = (
        "Curriculum Vitae profesional"
    )

    propiedades.keywords = (
        "CV, curriculum, candidatura"
    )


# ============================================================
# EXPORTACION
# ============================================================

def exportar_docx(
    cv: Dict[str, Any],
) -> Path:

    CARPETA_DOCUMENTOS.mkdir(
        parents=True,
        exist_ok=True,
    )

    documento = Document()

    configurar_documento(
        documento
    )

    configurar_metadatos(
        documento,
        cv,
    )

    agregar_cabecera(
        documento,
        cv,
    )

    agregar_perfil(
        documento,
        cv,
    )

    agregar_competencias(
        documento,
        cv,
    )

    agregar_experiencias(
        documento,
        cv,
    )

    agregar_formacion(
        documento,
        cv,
    )

    empresa = slug(
        cv.get(
            "vacante",
            {},
        ).get(
            "empresa",
            "EMPRESA",
        )
    )

    ruta = (
        CARPETA_DOCUMENTOS
        / f"CV_{empresa}_M_MONTANO.docx"
    )

    documento.save(
        ruta
    )

    return ruta


# ============================================================
# EJECUCION
# ============================================================

def main() -> None:

    vacante = cargar_vacante()

    cv = construir_cv(
        vacante
    )

    texto_control = renderizar_cv(
        cv
    )

    auditoria = auditar_cv(
        cv,
        texto_control,
    )

    if auditoria.get(
        "estado"
    ) != "APROBADO":

        print(
            "EXPORTACION BLOQUEADA"
        )

        print(
            "Incidencias:",
            auditoria.get(
                "incidencias",
                [],
            ),
        )

        return

    ruta = exportar_docx(
        cv
    )

    print(
        "DOCX GENERADO CORRECTAMENTE"
    )

    print(
        "Archivo:",
        ruta,
    )

    print(
        "Auditoria:",
        auditoria[
            "estado"
        ],
    )


if __name__ == "__main__":
    main()